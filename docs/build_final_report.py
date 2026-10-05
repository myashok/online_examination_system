from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    "/Users/user/Library/Application Support/Cursor/AgentStores/"
    "cursor_agent_stores/bc-e1643669-6440-465e-8127-a745bfede7da/files/"
    "docs/online-examination-system-mcsp-232-report.docx"
)
OUTPUT = ROOT / "docs" / "ONLINE-EXAMINATION-SYSTEM-MCSP-232-FINAL.docx"
ASSETS = ROOT / "docs" / "report-assets" / "current"


IMAGE_MAP = {
    "rId183": "login.png",
    "rId186": "student-registration.png",
    "rId189": "student-dashboard.png",
    "rId192": "exam-room.png",
    "rId195": "submission-confirmation.png",
    "rId198": "faculty-exam-list.png",
    "rId201": "question-bank.png",
    "rId204": "faculty-exam-editor.png",
    "rId207": "faculty-evaluation.png",
    "rId210": "faculty-ready-to-publish.png",
    "rId213": "admin-users.png",
    "rId383": "student-result-history.png",
    "rId386": "student-result-detail.png",
    "rId390": "faculty-report.png",
    "rId393": "faculty-analytics.png",
    "rId397": "admin-summary.png",
    "rId400": "admin-summary-detail.png",
    "rId404": "admin-audit.png",
    "rId407": "faculty-ready-to-publish.png",
}


REPLACEMENTS = {
    "Approved Proforma and Guide Bio-data": "Approved Proforma, Project Proposal and Guide Bio-data",
    "Attach the original approved project-proposal proforma and signed guide bio-data at this location before binding.": (
        "Attach (1) the original approved project-proposal proforma, (2) the approved project "
        "proposal/synopsis, and (3) the signed guide bio-data at this location before binding."
    ),
    "Appendix C — CD / USB handover": "Appendix C — CD handover",
    "Year of Submission: [TO BE PROVIDED]": "Year of Submission: 2026",
    "Study centre: [TO BE PROVIDED]": "Study centre: __________________________",
    "Regional centre: [TO BE PROVIDED]": "Regional centre: __________________________",
    "Guide qualification / designation / organisation: [TO BE PROVIDED]": (
        "Guide qualification / designation / organisation: ________________________________"
    ),
    "Project Proposal Number / approval date / evaluator stamp: [TO BE PROVIDED]": (
        "Project Proposal Number / approval date / evaluator stamp: ________________________"
    ),
    "Date: [TO BE PROVIDED]": "Date: __________________",
    "Address: [TO BE PROVIDED]": "Address: ______________________________________________",
    "Designation: [TO BE PROVIDED]": "Designation: ______________________________",
    "[TO BE PROVIDED]": "______________________________________________",
    "33 passed pytest cases in 5.61 seconds": "35 passed pytest cases in 6.57 seconds",
    "33 passed automated cases": "35 passed automated cases",
    "33 passed in 5.61s": "35 passed in 6.57s",
    "recorded 33-case run": "recorded 35-case run",
    "33 dots and 33 passed": "35 dots and 35 passed",
    "(33 passed)": "(35 passed)",
    "Executed automated cases (33 collected": "Executed automated cases (35 collected",
    ".................................": "...................................",
    "Figure 4.17 — Exam room: frozen paper, option order, palette, server timer.": (
        "Figure 4.17 — Exam room: frozen paper, server timer, palette filters and mark-for-review state."
    ),
    "Figure 4.20 — Question bank with MCQ/descriptive rows.": (
        "Figure 4.20 — Question bank with keyword, topic, difficulty and type filters."
    ),
    "Figure 4.21 — Exam schedule and paper controls.": (
        "Figure 4.21 — Published paper and lifecycle controls."
    ),
}


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_text(document):
    for paragraph in all_paragraphs(document):
        text = paragraph.text
        changed = text
        for old, new in REPLACEMENTS.items():
            changed = changed.replace(old, new)
        if changed != text:
            paragraph.text = changed


def paragraph(document, exact):
    return next(item for item in document.paragraphs if item.text == exact)


def update_project_specific_text(document):
    paragraph(document, "In scope (implemented). English-language browser UI; Student self-registration; admin-managed Faculty and Admin users; Argon2 hashing with PBKDF2 still accepted; SimpleJWT; subjects, topics, questions; exactly four MCQ options with one correct option; exam draft/publish/close; fixed paper and rule-based draw; one attempt per student per exam; frozen question order and frozen MCQ option order; server remaining-seconds; PUT auto-save; submit; auto-submit on expiry; negative marking floored at zero; unpublished then published results; query-based reports; admin audit list; oes_migrator versus oes_app.").text = (
        "In scope (implemented). English-language browser UI; Student self-registration; "
        "admin-managed Faculty and Admin users; Argon2 hashing with PBKDF2 still accepted; "
        "SimpleJWT; subjects, topics and questions; exactly four MCQ options with one correct "
        "option; keyword/topic/difficulty/type question-bank filters; exam draft/publish/close; "
        "fixed paper and rule-based draw; one attempt per student per exam; frozen question and "
        "option order; server remaining-seconds; auto-save; answered/not-answered/marked palette "
        "filters; persistent mark for review; submit and expiry auto-submit; negative marking "
        "floored at zero; unpublished then published results; faculty/admin CSV exports; "
        "query-based reports; admin audit list; oes_migrator versus oes_app."
    )
    paragraph(document, "The unique (attempt_id, question_id) row stores either a selected option or descriptive text. Award, evaluator and evaluation time are null until scoring. The selected option is protected so a historic answer does not lose its label.").text = (
        "The unique (attempt_id, question_id) row stores either a selected option or descriptive "
        "text plus marked_for_review. The review flag defaults false, persists independently of "
        "the answer and never changes scoring. Award, evaluator and evaluation time are null "
        "until scoring. The selected option is protected so a historic answer does not lose its label."
    )
    paragraph(document, "Inputs. Authenticated Student and exam id at start; attempt id at detail/save/submit; answer list during save.").text = (
        "Inputs. Authenticated Student and exam id at start; attempt id at detail/save/submit; "
        "answer list during save with selected_option_id, descriptive_text and marked_for_review."
    )
    paragraph(document, "Processing. Start checks actor, status and window, creates one attempt, shuffles and freezes paper/options. Detail calls ensure_open and calculates remaining time. Save validates ownership, status, paper membership and option membership. Submit locks the row, records time, scores objectives, and finalizes only if descriptive work is complete.").text = (
        "Processing. Start checks actor, status and window, creates one attempt, shuffles and "
        "freezes paper/options. Detail calls ensure_open and calculates remaining time. Save "
        "validates ownership, status, paper membership and option membership, then persists the "
        "answer and optional review flag. Submit locks the row, records time, scores objectives, "
        "and finalizes only if descriptive work is complete."
    )
    paragraph(document, "Outputs. The attempt payload includes exam metadata, status, deadline, server time, remaining seconds and ordered questions. During the sitting, option objects contain id/text only. Submit returns status/result-ready and publish state.").text = (
        "Outputs. The attempt payload includes exam metadata, status, deadline, server time, "
        "remaining seconds, ordered questions and marked_for_review for each item. During the "
        "sitting, option objects contain id/text only. Submit returns status/result-ready and publish state."
    )
    paragraph(document, "MCQ scoring — blank 0; correct = marks; wrong = − negative_mark.").text = (
        "Current save and bank-filter excerpts — review is independent of scoring; search is "
        "case-insensitive and combines with structured filters."
    )
    paragraph(document, "if answer.selected_option_id is None:\n    answer.awarded_marks = Decimal(\"0.00\")\nelif answer.selected_option.is_correct:\n    answer.awarded_marks = slot.question.marks\nelse:\n    answer.awarded_marks = Decimal(\"0.00\") - negative").text = (
        'if "marked_for_review" in item and item.get("marked_for_review") is not None:\n'
        "    answer.marked_for_review = bool(item.get(\"marked_for_review\"))\n"
        "answer.save()\n\n"
        'search = (request.query_params.get("q") or request.query_params.get("search") or "").strip()\n'
        "if search:\n"
        "    questions = questions.filter(text__icontains=search)"
    )
    paragraph(document, "Browser workflow (separate from pytest): register/login, start, save, submit, mark, publish, result, reports, audit.").text = (
        "Browser workflow (separate from pytest): register/login, bank search/filter, start, "
        "save, mark/unmark for review, submit, descriptive mark, publish, result, CSV/report and audit."
    )
    paragraph(document, "Dedicated pytest coverage is also absent for token refresh, subject/topic CRUD, exam unpublish, and several list/detail APIs (honest gap).").text = (
        "Dedicated pytest coverage remains absent for token refresh, subject/topic CRUD, exam "
        "unpublish and several list/detail APIs. The final 35-case run adds persistence coverage "
        "for mark-for-review and combined question-bank filtering."
    )
    paragraph(document, "Dated 5 October 2026. The test database was created with the migrator role. Privilege tests connected as oes_app and confirmed CREATE TABLE was denied. This snapshot is not a promise that an altered future build will produce the same result.").text = (
        "Dated 5 October 2026. The test database was created with the migrator role. Privilege "
        "tests connected as oes_app and confirmed CREATE TABLE was denied. The run includes the "
        "review-flag and question-bank-filter regression cases. This snapshot applies to the final source package."
    )


def find_table(document, first_headers, contains=None):
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers[: len(first_headers)] != first_headers:
            continue
        if contains and not any(contains in cell.text for row in table.rows for cell in row.cells):
            continue
        return table
    raise LookupError((first_headers, contains))


def add_row(table, values):
    cells = table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def update_tables(document):
    requirements = find_table(document, ["ID", "Shall statement", "Actor", "Verification"])
    add_row(requirements, ["FR-40", "Save, return and clear a per-question mark-for-review flag without affecting marks", "Student", "INT-COND-07"])
    add_row(requirements, ["FR-41", "Search/filter the question bank by keyword, topic, difficulty and type", "Faculty, Admin", "INT-QB-06"])
    add_row(requirements, ["FR-42", "Export faculty exam and admin institution summaries as CSV", "Faculty, Admin", "SYS-RPT-01"])

    answers = find_table(document, ["Column", "Type", "Null", "Key"], "descriptive_text")
    add_row(answers, ["marked_for_review", "tinyint(1)", "N", "DEFAULT 0"])

    ui = find_table(document, ["Screen", "Route", "Principal controls"])
    for row in ui.rows[1:]:
        name = row.cells[0].text.strip().lower()
        if "question bank" in name:
            row.cells[2].text = "Keyword, topic, difficulty and type filters; result count; new question"
        if "exam room" in name:
            row.cells[2].text = "Server timer; answered/not-answered/marked filters; palette; mark/unmark; save; submit"

    tests = find_table(document, ["Test ID", "Case", "Input / condition", "Expected", "Actual", "Result"])
    add_row(tests, ["INT-COND-07", "mark for review", "save true; reload; clear false", "flag persists; answer/scoring unchanged", "matched", "PASS"])
    add_row(tests, ["INT-QB-06", "question-bank filters", "keyword/topic/difficulty/type combinations", "only matching questions returned", "matched", "PASS"])

    trace = find_table(document, ["Requirement", "Module", "Table(s)", "API", "Test", "Report"])
    add_row(trace, ["FR-40", "Conduction", "attempt_answers", "PUT /api/attempts/{id}/answers/", "INT-COND-07", "Figures 4.17, 6.3"])
    add_row(trace, ["FR-41", "Question bank", "questions/topics", "GET /api/questions/?q=&topic=&difficulty=&type=", "INT-QB-06", "Figure 4.20"])
    add_row(trace, ["FR-42", "Reports", "query output", "GET report export endpoints", "SYS-RPT-01", "Figures 8.3–8.6"])


def replace_screenshots(document):
    for shape in document.inline_shapes:
        blips = shape._inline.xpath('.//*[local-name()="blip"]')
        if not blips:
            continue
        rid = blips[0].get(qn("r:embed"))
        filename = IMAGE_MAP.get(rid)
        if not filename:
            continue
        image = ASSETS / filename
        if not image.exists():
            raise FileNotFoundError(image)
        document.part.rels[rid].target_part._blob = image.read_bytes()
        shape.width = Inches(6.22)
        shape.height = Inches(4.32)


def main():
    document = Document(SOURCE)
    replace_text(document)
    update_project_specific_text(document)
    update_tables(document)
    replace_screenshots(document)
    document.core_properties.title = "ONLINE EXAMINATION SYSTEM — MCSP-232 Final Project Report"
    document.core_properties.subject = "Final project documentation checked against MCSP-232 guidelines"
    document.core_properties.comments = (
        "Technical final prepared 5 October 2026. Signatures, approved proforma, guide biodata "
        "and regional-centre particulars must be completed before physical submission."
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
